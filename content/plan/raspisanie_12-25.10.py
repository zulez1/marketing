"""Расписание 12–25.10 по площадкам: VK + Instagram — привлечение и лиды, Telegram + чат — комьюнити и важная информация.
python3 raspisanie_12-25.10.py → raspisanie-12-25.10.xlsx"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter

INK, AMBER, PAPER = "0D0F13", "F0B90B", "F3F0E8"
F = lambda **k: Font(name="Arial", **k)
thin = Side(style="thin", color="DAD5C6"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
N, D, S = "Нужно", "Выложено", "—"
W = "content/nedeli-3-4_12.10-25.10/"

# дата, день, время, тип, тема, Instagram, VK, Telegram-канал, Telegram-чат, сторис (IG+VK), папка
rows = [
 ("12.10", "Пн", "12:00 / 18:00", "Знакомство · форматы", "Знакомьтесь: зубр + 9 форматов форума",
  "Рилс «Знакомьтесь, это зубр» (12:00, версия без звука + трендовый трек, удар на 2-й секунде) + карусель 5 сл. (18:00) · CTA: подпишитесь", "Клип с зубром + карусель · CTA: подпишитесь, ссылка на сайт",
  "Пост «9 форматов форума» полностью + рилс", "Знакомство: «Кто вы, какой бизнес, какой город?» — первое сообщение недели",
  "Репост рилса + опрос «Какой формат интереснее?» · отсчёт 32", "01-pn-12.10-zubr-formaty"),
 ("13.10", "Вт", "10:00", "Польза · лид-магнит", "Сколько стоит потерянный клиент + чек-лист «12 причин» (PDF)",
  "Карусель 6 сл. · CTA: слово КЛИЕНТЫ → в директ ссылка на чек-лист в TG", "Карусель + текст · CTA: чек-лист в Telegram (ссылка)",
  "Пост + PDF-чек-лист «12 причин» отдельным сообщением (закрепить на день)", "«Кто считал LTV? Сколько месяцев в среднем ходит ваш клиент?»",
  "Зубр в шоке + викторина «Сколько ходит клиент» · отсчёт 31", "02-vt-13.10-poteryannyy-klient"),
 ("14.10", "Ср", "12:00 + 19:00", "Розыгрыш + мем", "Розыгрыш билета (итоги 25.10) + мем «Родительский чат»",
  "12:00 розыгрыш (2 сл.) · 19:00 мем · CTA: отметьте коллегу", "Розыгрыш + мем · CTA: подписка и отметка коллеги", "Пост с полными условиями розыгрыша (закрепить) + мем",
  "«Во сколько вам пришло последнее сообщение от родителя?» + ссылка на розыгрыш", "Сторис «Разыгрываем билет» + стикер-ссылка · отсчёт 30", "03-sr-14.10-rozygrysh-bileta"),
 ("15.10", "Чт", "10:00", "Польза", "5 фраз, которые удерживают клиента после пробного занятия",
  "Карусель 7 сл. · CTA: сохраните, отправьте админу", "Карусель + текст · CTA: сохраните", "Пост с фразами полностью + ссылка на сайт",
  "«Какая фраза у вас лучше всего удерживает после пробного?»", "Обложка + «Сохраните» · отсчёт 29", "04-cht-15.10-5-fraz-posle-probnogo"),
 ("16.10", "Пт", "19:00", "Мем + рилс", "«Нет / Да» про форумы (верх — ч/б зубр) + рилс «Полный зал»",
  "Мем + рилс под трендовый звук · CTA: ссылка в профиле", "Мем + клип · CTA: ссылка на сайт", "Мем + рилс",
  "«Пятничный вопрос: лучший и худший форум, на котором вы были?»", "Мем + опрос «Какой форум выберете?» · напоминание о розыгрыше · отсчёт 28", "05-pt-16.10-mem-net-da"),
 ("17.10", "Сб", "12:00", "Форум", "Как выглядит день на РЕШЕНО (13.11 по часам)",
  "Карусель 6 сл. · CTA: программа по ссылке в профиле", "Карусель + текст · CTA: программа на сайте (#program)", "Пост с полной программой дня — важная информация, закрепить",
  "«Какую сессию 13 ноября ждёте больше всего?»", "Слайды «Утро» и «Вечер» + вопрос · отсчёт 27", "06-sb-17.10-den-na-reshno"),
 ("18.10", "Вс", "12:00", "Вовлечение", "«Какой ты зубр сегодня?»",
  "Картинка «Какой ты зубр» · CTA: цифра в комментариях", "Картинка · CTA: цифра в комментариях", "«Какой ты зубр» + напоминание о розыгрыше",
  "«Какой вы зубр сегодня?»", "Картинка + опрос · отсчёт 26", "07-vs-18.10-kakoy-ty-zubr"),
 ("19.10", "Пн", "10:00", "Польза", "6 цифр, которые владелец смотрит каждый понедельник",
  "Карусель 5 сл. · CTA: сохраните", "Карусель + текст · CTA: сессия Finance на сайте", "Пост с цифрами полностью",
  "«Покажите, какие цифры смотрите вы — таблица, CRM, тетрадка?»", "Флипчарт + шкала «Сколько цифр из 6 считаете» · отсчёт 25", "08-pn-19.10-6-cifr-ponedelnika"),
 ("20.10", "Вт", "19:00", "Мем", "Ожидание / реальность: второй филиал",
  "Мем · CTA: отметьте того, кто собирается открывать", "Мем · CTA: отметьте коллегу", "Мем + анонс дебатов Owner vs Owner",
  "«У кого 2+ филиала — главная ошибка при открытии второго?»", "Мем + опрос «Второй филиал: есть / думаю / ни за что» · отсчёт 24", "09-vt-20.10-mem-ozhidanie-realnost"),
 ("21.10", "Ср", "10:00", "Польза", "5 вопросов тренеру на собеседовании",
  "Карусель 7 сл. · CTA: сохраните к найму", "Карусель + текст · CTA: сессия Team на сайте", "Пост с вопросами полностью",
  "«Ваш главный вопрос на собеседовании тренера?»", "Обложка + стикер «Вопрос» · отсчёт 23", "10-sr-21.10-sobesedovanie-trenera"),
 ("22.10", "Чт", "19:00", "Мем + рилс", "Уровни владельца + рилс «Hot Seat»",
  "Мем + рилс под трендовый звук · CTA: ссылка в профиле", "Мем + клип · CTA: ссылка на сайт", "Мем + рилс",
  "«Какой вопрос вы бы задали на Hot Seat, если бы никто не узнал, что это вы?»", "Мем + шкала уровней · репост рилса · отсчёт 22", "11-cht-22.10-mem-urovni"),
 ("23.10", "Пт", "10:00", "Форум · заявки", "Business Live: подайте бизнес на разбор",
  "Карусель 4 сл. · CTA: заявка по ссылке в профиле", "Карусель + текст · CTA: заявка на сайте (#business-live)", "Пост: как работает разбор, даты, ссылка на заявку — закрепить",
  "«Кто подаёт бизнес на разбор? Вопросы по заявке — сюда»", "Обложка + стикер-ссылка на заявку · отсчёт 21", "12-pt-23.10-business-live"),
 ("24.10", "Сб", "12:00", "Продажа", "Вы покупаете не билет + стартовая цена для первых 50",
  "Карусель 4 сл. · CTA: забронировать (ссылка в профиле)", "Карусель + текст · CTA: забронировать (#tickets)", "Пост: что входит, пакеты 3 и 5, стартовая цена — важная информация",
  "«Едете командой? Ищем, с кем объединиться в пакет на 3 или 5 человек»", "Зубр + стикер-ссылка «Забронировать» · отсчёт 20", "13-sb-24.10-ne-bilet"),
 ("25.10", "Вс", "12:00", "Итоги розыгрыша", "Итоги розыгрыша билета",
  "Картинка-победитель · CTA: стартовая цена для первых 50", "Картинка · CTA: забронировать (#tickets)", "Пост-итоги + видео выбора победителя",
  "Поздравляем победителя + итоги двух недель", "Сторис-победитель + отметка · отсчёт 19", "14-vs-25.10-itogi-rozygrysha"),
]
PL = ["Instagram", "VK", "Telegram-канал", "Telegram-чат", "Сторис IG+VK"]

wb = Workbook()
# ---------- лист 1: расписание ----------
ws = wb.active; ws.title = "Расписание 12–25.10"
head = ["Дата", "День", "Время", "Тип", "Тема"] + [f"{p}: что выкладываем" for p in PL] + [f"{p}: статус" for p in PL] + ["Готово?", "Папка с материалами"]
ncol = len(head)
ws.append(["РЕШЕНО · расписание публикаций 12–25 октября"]); ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
ws["A1"].font = F(bold=True, size=16, color=AMBER); ws["A1"].fill = PatternFill("solid", fgColor=INK); ws["A1"].alignment = Alignment(vertical="center", indent=1); ws.row_dimensions[1].height = 32
ws.append(["VK и Instagram — привлекаем новых людей и лиды (охват, мемы, рилсы, ссылка на сайт и в Telegram). Telegram-канал — главная площадка: полные тексты, программа, условия, чек-листы. Telegram-чат — живое общение: вопрос дня, нетворкинг. В колонках «статус» выберите «Выложено», когда опубликовали."])
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol); ws["A2"].font = F(italic=True, size=10, color="585D66"); ws["A2"].alignment = Alignment(wrap_text=True, vertical="center"); ws.row_dimensions[2].height = 34
ws.append(head); HR = 3
groupfill = {"Instagram": "E9D7F5", "VK": "D6E4F7", "Telegram-канал": "D4EEF7", "Telegram-чат": "D9F0E2", "Сторис IG+VK": "F7E7C6"}
for c in range(1, ncol + 1):
    cell = ws.cell(HR, c); cell.font = F(bold=True, color="FFFFFF", size=10); cell.fill = PatternFill("solid", fgColor=INK)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = B
ws.row_dimensions[HR].height = 34
first = HR + 1
SC0 = 6 + len(PL)  # first status column
for i, r in enumerate(rows):
    rr = first + i
    date, day, time, typ, topic, *cont, folder = r
    sc1, sc2 = get_column_letter(SC0), get_column_letter(SC0 + len(PL) - 1)
    ready = f'=IF(COUNTIF({sc1}{rr}:{sc2}{rr},"{N}")=0,"Всё выложено","Осталось: "&COUNTIF({sc1}{rr}:{sc2}{rr},"{N}"))'
    ws.append([date, day, time, typ, topic] + cont + [N] * len(PL) + [ready, W + folder])
    for c in range(1, ncol + 1):
        cell = ws.cell(rr, c); cell.border = B; cell.font = F(size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if c <= 3 or c >= SC0 and c < SC0 + len(PL) + 1 else "left")
    ws.cell(rr, 5).font = F(size=10, bold=True)
    for k, p in enumerate(PL):
        ws.cell(rr, 6 + k).fill = PatternFill("solid", fgColor=groupfill[p])
    if day in ("Сб", "Вс"):
        for c in (1, 2, 3): ws.cell(rr, c).fill = PatternFill("solid", fgColor="F3F0E8")
    ws.row_dimensions[rr].height = 78
last = first + len(rows) - 1
rng = f"{get_column_letter(SC0)}{first}:{get_column_letter(SC0 + len(PL) - 1)}{last}"
dv = DataValidation(type="list", formula1=f'"{N},{D},{S}"', allow_blank=True); ws.add_data_validation(dv); dv.add(rng)
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{D}"'], fill=PatternFill("solid", fgColor="CDEBD6"), font=Font(name="Arial", color="1E6B3A", bold=True)))
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{N}"'], fill=PatternFill("solid", fgColor="FCEFC2"), font=Font(name="Arial", color="7A5B00")))
rc = get_column_letter(SC0 + len(PL))
ws.conditional_formatting.add(f"{rc}{first}:{rc}{last}", CellIsRule(operator="equal", formula=['"Всё выложено"'], fill=PatternFill("solid", fgColor="CDEBD6"), font=Font(name="Arial", color="1E6B3A", bold=True)))
widths = [7, 5, 8, 16, 30] + [30] * len(PL) + [12] * len(PL) + [14, 40]
for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = ws.cell(first, 6); ws.auto_filter.ref = f"A{HR}:{get_column_letter(ncol)}{last}"

# ---------- лист 2: роли площадок ----------
r2 = wb.create_sheet("Роли площадок")
roles = [
 ("Площадка", "Роль", "Что выкладываем", "Призыв (CTA)", "Как мерить"),
 ("Instagram", "Привлечение новых людей и лидов", "Карусели-полезняшки, мемы с зубром, рилсы под трендовый звук, сторис каждый день", "Ссылка в профиле на сайт · слово в комментариях → лид-магнит в Telegram · соавторство со спикерами", "Охват, сохранения, переходы по ссылке, заявки с utm_source=instagram"),
 ("VK", "Привлечение новых людей и лидов", "Те же карусели и мемы, клипы, сторис", "Ссылка на сайт (utm_source=vk) · ссылка на Telegram-канал", "Охват, переходы, заявки с utm_source=vk"),
 ("Telegram-канал", "Главная площадка комьюнити и важной информации", "Полные тексты, программа, условия конкурсов, чек-листы, анонсы, опросы", "Ссылка на сайт (utm_source=telegram) · закреп с главным", "Подписчики, просмотры, реакции, заявки с utm_source=telegram"),
 ("Telegram-чат", "Живое общение участников", "Вопрос дня, знакомства, обсуждения, поиск партнёров и попутчиков на форум", "Отвечать, знакомить людей между собой, звать в канал за подробностями", "Сообщения в день, новые участники, кто представился"),
 ("Сторис IG + VK", "Ежедневный контакт", "Отсчёт «До форума N дней», репост поста дня, опросы и вопросы", "Стикер-ссылка на сайт или пост", "Ответы на стикеры, переходы"),
]
for i, row in enumerate(roles):
    r2.append(row)
    for c in range(1, 6):
        cell = r2.cell(i + 1, c); cell.border = B; cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = F(bold=True, color="FFFFFF") if i == 0 else F(size=10, bold=(c == 1))
        if i == 0: cell.fill = PatternFill("solid", fgColor=INK)
    r2.row_dimensions[i + 1].height = 22 if i == 0 else 64
for i, w in enumerate([18, 30, 46, 46, 40], 1): r2.column_dimensions[get_column_letter(i)].width = w

# ---------- лист 3: сводка ----------
s = wb.create_sheet("Сводка")
s.append(["Площадка", "Запланировано", "Выложено", "Осталось", "Прогресс"])
for c in range(1, 6): s.cell(1, c).font = F(bold=True, color="FFFFFF"); s.cell(1, c).fill = PatternFill("solid", fgColor=INK)
for k, p in enumerate(PL):
    col = get_column_letter(SC0 + k); R = f"'Расписание 12–25.10'!{col}{first}:{col}{last}"; r = 2 + k
    s.append([p, f'=COUNTIF({R},"{N}")+COUNTIF({R},"{D}")', f'=COUNTIF({R},"{D}")', f"=B{r}-C{r}", f"=IF(B{r}=0,0,C{r}/B{r})"])
    s.cell(r, 5).number_format = "0%"
for i, w in enumerate([18, 15, 12, 12, 12], 1): s.column_dimensions[get_column_letter(i)].width = w
wb.calculation.fullCalcOnLoad = True
wb.save("raspisanie-12-25.10.xlsx"); print("ok", len(rows))
